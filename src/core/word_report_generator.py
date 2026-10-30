"""Generate Word report from structured JSON output (LLM response).

Usage:
    from core.word_report_generator import generate_word_report, translate_brain_region
    word_bytes = generate_word_report(report_json)
"""
from __future__ import annotations

import io
import os
import re
from functools import lru_cache
from pathlib import Path
from typing import Any, Mapping

import yaml

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ── Brain region translation table ──────────────────────────────
BRAIN_REGION_ZH: dict[str, str] = {
    "prefrontal cortex": "前额叶皮层",
    "motor cortex": "运动皮层",
    "parietal cortex": "顶叶皮层",
    "superior temporal gyrus": "颞上回",
    "middle temporal gyrus": "颞中回",
    "inferior frontal gyrus": "额下回",
    "superior frontal gyrus": "额上回",
    "temporoparietal junction": "颞顶联合区",
    "inferior parietal lobule": "顶下小叶",
    "cingulate cortex": "扣带回",
    "anterior cingulate cortex": "前扣带回",
    "insula": "岛叶",
    "amygdala": "杏仁核",
    "basal ganglia": "基底节",
    "occipital cortex": "枕叶皮层",
    "dorsolateral prefrontal cortex": "背外侧前额叶",
    "ventrolateral prefrontal cortex": "腹外侧前额叶",
    "supplementary motor area": "辅助运动区",
    "premotor cortex": "前运动皮层",
    "posterior parietal cortex": "后顶叶皮层",
    "hippocampus": "海马体",
    "thalamus": "丘脑",
    "cerebellum": "小脑",
    "whole-head": "全头覆盖",
}


def translate_brain_region(en_name: str) -> str:
    """Translate a brain region name from English to Chinese."""
    return BRAIN_REGION_ZH.get(en_name.strip().lower(), en_name)


# ─ Markdown cleanup ───────────────────────────────────────────

def _clean_markdown(text: str) -> str:
    """Remove markdown formatting characters from text.

    Converts markdown syntax to plain text for Word rendering.
    """
    if not text:
        return text

    # Remove markdown link syntax [text](url) → text
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)

    # Remove markdown heading markers (# ## ### etc.)
    text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)

    # Remove bold/italic markers but keep the text
    text = re.sub(r'\*\*([^\*]+)\*\*', r'\1', text)  # **bold**
    text = re.sub(r'\*([^\*]+)\*', r'\1', text)      # *italic*
    text = re.sub(r'__([^\_]+)__', r'\1', text)      # __bold__
    text = re.sub(r'_([^\_]+)_', r'\1', text)        # _italic_

    # Remove list markers at start of line
    text = re.sub(r'^[-*]\s+', '', text, flags=re.MULTILINE)
    text = re.sub(r'^\d+\.\s+', '', text, flags=re.MULTILINE)

    # Remove code block markers
    text = re.sub(r'`{1,3}([^\`]+)`{1,3}', r'\1', text)

    # Remove blockquote markers
    text = re.sub(r'^>\s+', '', text, flags=re.MULTILINE)

    # Remove horizontal rule markers
    text = re.sub(r'^[-*_]{3,}\s*$', '', text, flags=re.MULTILINE)

    # Clean up extra whitespace
    text = re.sub(r'\n{3,}', '\n\n', text)  # Max 2 consecutive newlines
    text = text.strip()

    return text


# ── Default report theme (matches the app's blue theme) ───────────────────────────────────────
_DEFAULT_THEME_HEX = {
    "primary": "005387",
    "secondary": "00375B",
    "accent": "1479B0",
    "gray": "627482",
    "light_gray": "EFF6FB",
}
_DEFAULT_EXPORT_CONFIG_PATH = Path(__file__).resolve().parents[2] / "config" / "export.yaml"


def _rgb_from_hex(value: str, fallback: str) -> RGBColor:
    normalized = value.strip().removeprefix("#") if isinstance(value, str) else ""
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", normalized):
        normalized = fallback
    return RGBColor.from_string(normalized.upper())


@lru_cache(maxsize=16)
def _load_export_config(path: str, modified_ns: int) -> dict[str, Any]:
    del modified_ns
    try:
        with open(path, encoding="utf-8") as stream:
            loaded = yaml.safe_load(stream) or {}
    except (OSError, yaml.YAMLError):
        return {}
    return loaded if isinstance(loaded, dict) else {}


def load_word_theme(config_path: str | os.PathLike[str] | None = None) -> dict[str, RGBColor]:
    """Load cached Word report colors from the export config or defaults."""
    configured_path = config_path or os.environ.get("EXPORT_CONFIG")
    path = Path(configured_path).expanduser() if configured_path else _DEFAULT_EXPORT_CONFIG_PATH
    try:
        resolved = path.resolve()
        loaded = _load_export_config(str(resolved), resolved.stat().st_mtime_ns)
    except OSError:
        loaded = {}
    word = loaded.get("word") if isinstance(loaded.get("word"), dict) else {}
    configured_theme = word.get("theme") if isinstance(word.get("theme"), dict) else {}
    return {
        name: _rgb_from_hex(configured_theme.get(name, fallback), fallback)
        for name, fallback in _DEFAULT_THEME_HEX.items()
    }


THEME_PRIMARY = RGBColor.from_string(_DEFAULT_THEME_HEX["primary"])
THEME_SECONDARY = RGBColor.from_string(_DEFAULT_THEME_HEX["secondary"])
THEME_ACCENT = RGBColor.from_string(_DEFAULT_THEME_HEX["accent"])
THEME_GRAY = RGBColor.from_string(_DEFAULT_THEME_HEX["gray"])
THEME_LIGHT_GRAY = RGBColor.from_string(_DEFAULT_THEME_HEX["light_gray"])

# ── Style helpers ───────────────────────────────────────────────

def _set_cell_shading(cell: Any, color: str) -> None:
    """Set cell background color."""
    sh = cell._element.get_or_add_tcPr()
    bg = sh.makeelement(qn("w:shd"), {qn("w:fill"): color, qn("w:val"): "clear"})
    sh.append(bg)


def _setup_styles(doc: Document, theme: Mapping[str, RGBColor]) -> None:
    """Configure document styles (fonts, heading sizes, colors)."""
    # Normal style
    style = doc.styles["Normal"]
    style.font.name = "思源黑体 CN Normal"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.makeelement(qn("w:rFonts"), {})
    rFonts.set(qn("w:eastAsia"), "思源黑体 CN Normal")
    rPr.insert(0, rFonts)

    # Heading styles
    for level, sz in [(1, 16), (2, 13)]:
        hs = doc.styles[f"Heading {level}"]
        hs.font.name = "思源黑体 CN Normal"
        hs.font.size = Pt(sz)
        hs.font.bold = True
        hs.font.color.rgb = theme["primary"]
        hs.paragraph_format.space_before = Pt(18 if level == 1 else 14)
        hs.paragraph_format.space_after = Pt(8)


def _add_page_numbers(doc: Document, theme: Mapping[str, RGBColor]) -> None:
    """Add a centered dynamic PAGE field to every distinct footer."""
    seen_footers: set[int] = set()
    for section in doc.sections:
        footer = section.footer
        footer_id = id(footer._element)
        if footer_id in seen_footers:
            continue
        seen_footers.add(footer_id)
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.font.size = Pt(9)
        run.font.color.rgb = theme["gray"]

        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        placeholder = OxmlElement("w:t")
        placeholder.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        for element in (begin, instruction, separate, placeholder, end):
            run._r.append(element)


# ── Content rendering ───────────────────────────────────────────

def _add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    theme: Mapping[str, RGBColor],
) -> None:
    """Render a table with styled header and alternating rows."""
    t = doc.add_table(rows=1 + len(rows), cols=len(headers), style="Table Grid")
    t.autofit = True

    # Header row uses the configured primary theme color.
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = ""
        r = c.paragraphs[0].add_run(_clean_markdown(h))
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cell_shading(c, "".join(f"{channel:02X}" for channel in theme["primary"]))

    # Data rows
    for i, row_data in enumerate(rows):
        for j, val in enumerate(row_data):
            c = t.rows[i + 1].cells[j]
            c.text = ""
            r = c.paragraphs[0].add_run(_clean_markdown(str(val) if val else ""))
            r.font.size = Pt(9)
            # Alternating row color
            if i % 2 == 0:
                _set_cell_shading(c, "".join(f"{channel:02X}" for channel in theme["light_gray"]))


def _add_bullets(doc: Document, items: list[str]) -> None:
    """Render a bullet list."""
    for item in items:
        p = doc.add_paragraph(_clean_markdown(str(item)), style="List Bullet")
        for run in p.runs:
            run.font.size = Pt(10)


def _add_references(doc: Document, references: list[dict[str, Any]]) -> None:
    """Render numbered APA references with hanging indents and italics."""
    for reference in references:
        number = int(reference.get("number") or 0)
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.75)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.keep_together = True
        prefix = paragraph.add_run(f"[{number}] ")
        prefix.font.size = Pt(10)
        segments = reference.get("segments")
        if isinstance(segments, list) and segments:
            for segment in segments:
                if not isinstance(segment, dict):
                    continue
                run = paragraph.add_run(_clean_markdown(str(segment.get("text") or "")))
                run.font.size = Pt(10)
                run.font.italic = bool(segment.get("italic"))
        else:
            run = paragraph.add_run(_clean_markdown(str(reference.get("text") or "")))
            run.font.size = Pt(10)


def _add_section(doc: Document, section: dict, theme: Mapping[str, RGBColor]) -> None:
    """Recursively render a section (heading + content + table + subsections + bullets)."""
    heading = section.get("heading", "")
    level = section.get("level", 1)

    if heading:
        doc.add_heading(_clean_markdown(heading), level=min(level, 2))

    # Content paragraph
    content = section.get("content", "")
    if content:
        doc.add_paragraph(_clean_markdown(content))

    # Table
    table_data = section.get("table")
    if table_data:
        headers = table_data.get("headers", [])
        rows = table_data.get("rows", [])
        if headers and rows:
            _add_table(doc, headers, rows, theme)

    # Also handle type=="table" with top-level headers/rows
    if section.get("type") == "table":
        headers = section.get("headers", [])
        rows = section.get("rows", [])
        if headers and rows:
            _add_table(doc, headers, rows, theme)

    # Subsections
    for sub in section.get("subsections", []):
        _add_section(doc, sub, theme)

    # Bullets
    bullets = section.get("bullets", [])
    if bullets:
        _add_bullets(doc, bullets)

    references = section.get("references", [])
    if references:
        _add_references(doc, references)


# ── Main entry ──────────────────────────────────────────────────

def generate_word_report(report_json: dict) -> bytes:
    """Generate a Word document from structured report JSON.

    Args:
        report_json: Dict with keys:
            - report_title (str)
            - report_subtitle (str, optional)
            - overview (str, optional)
            - sections (list of section dicts)

    Returns:
        bytes: The .docx file content.
    """
    doc = Document()
    theme = load_word_theme()
    _setup_styles(doc, theme)
    _add_page_numbers(doc, theme)

    # Title
    title_text = _clean_markdown(report_json.get("report_title", "研究报告"))
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(title_text)
    tr.font.size = Pt(22)
    tr.font.bold = True
    tr.font.color.rgb = theme["primary"]

    # Subtitle
    subtitle = _clean_markdown(report_json.get("report_subtitle", ""))
    if subtitle:
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sp.add_run(subtitle)
        sr.font.size = Pt(10)
        sr.font.color.rgb = theme["gray"]

    # Overview
    overview = _clean_markdown(report_json.get("overview", ""))
    if overview:
        doc.add_paragraph(overview)

    # Sections
    for section in report_json.get("sections", []):
        _add_section(doc, section, theme)

    # Serialize to bytes
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
