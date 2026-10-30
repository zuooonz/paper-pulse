"""预览 Word 样式修改效果"""
from docx import Document
from docx.shared import RGBColor, Pt
from docx.oxml.ns import qn

# 加载文档
doc = Document(r'C:\Users\zuooo\Downloads\论文总结_20260717 (1).docx')

# 定义新颜色（辅助蓝）
NEW_BLUE = RGBColor(62, 130, 185)  # R62 G130 B185
WHITE = RGBColor(255, 255, 255)

# 修改字体为思源黑体
def update_font(run, font_name='思源黑体'):
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.makeelement(qn('w:rFonts'), {})
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.insert(0, rFonts)

# 遍历所有段落
for para in doc.paragraphs:
    for run in para.runs:
        update_font(run)
        # 标题和标题样式使用新蓝色
        if run.bold and run.font.size and run.font.size >= Pt(13):
            run.font.color.rgb = NEW_BLUE

# 遍历所有表格
for table in doc.tables:
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    update_font(run)
                    if i == 0:  # 表头
                        run.font.color.rgb = WHITE
                        run.font.bold = True

# 保存预览文件
output_path = r'C:\Users\zuooo\Downloads\论文总结_20260717_preview2.docx'
doc.save(output_path)
print(f'预览文件已保存：{output_path}')
