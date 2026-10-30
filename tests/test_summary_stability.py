import json
from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from api import summary_jobs
from api.summary_service import (
    build_deterministic_report,
    build_reference_number_map,
    generate_report_json,
    validate_final_report,
    validate_narrative_payload,
)
from paper_library import db
from paper_library.governance_models import GovernanceBase
from paper_library.governance_models import SummaryJob
from paper_library.time_utils import utc_now
from core.word_report_generator import generate_word_report


def _papers():
    return [
        {
            "id": 1,
            "internal_id": 727,
            "title": "Paper A",
            "title_zh": "论文甲",
            "journal_name": "Journal A",
            "year": 2025,
            "impact_factor": 6.2,
            "jcr_quartile": "Q1",
            "population": ["儿童"],
            "paradigm": ["Go/No-Go"],
            "research_fields": ["执行控制"],
            "brain_regions": ["前额叶皮层"],
            "author_names": ["Author Alpha", "Author Beta"],
            "volume": "12",
            "issue": "3",
            "pages": "10-18",
            "doi": "10.1000/paper-a",
        },
        {
            "id": 2,
            "internal_id": 643,
            "title": "Paper B",
            "journal_name": "Journal B",
            "year": 2024,
            "impact_factor": 3.1,
            "jcr_quartile": "Q2",
            "population": ["成人"],
            "paradigm": ["静息态"],
            "research_fields": ["脑网络"],
            "brain_regions": [],
            "author_names": ["Author Gamma"],
            "doi": "10.1000/paper-b",
        },
    ]


def _payload():
    content = ("这些研究围绕共同问题形成了可比较的观察路径，并呈现出清晰的方法差异与结果联系。" * 2) + "〔REF_001〕〔REF_002〕"
    return {
        "report_title": "执行控制与脑网络研究",
        "report_subtitle": "qMRI 综合分析",
        "overview": "所选研究覆盖不同年龄群体和任务情境，研究设计从任务态测量延伸到静息态网络观察。各项工作在研究对象、范式安排和分析重点上形成互补，为理解相关脑功能变化提供了连贯视角。",
        "narrative": {
            "technical_scope": {"content": content},
            "population_topics": {
                "content": content,
                "subsections": [
                    {"heading": "儿童研究", "content": content},
                    {"heading": "成人研究", "content": content},
                ],
            },
            "experimental_design": {"content": content},
            "method_analysis": {"heading": "脑区覆盖与神经同步特征", "content": content},
            "core_findings": {
                "content": content,
                "subsections": [
                    {"heading": "任务相关变化", "content": content},
                    {"heading": "网络层面联系", "content": content},
                ],
            },
            "conclusion": {
                "content": content,
                "bullets": ["要点一", "要点二", "要点三", "要点四", "要点五"],
            },
        },
    }


def test_program_owns_seven_sections_without_tables_and_with_references():
    report = build_deterministic_report(_payload(), _papers())

    assert len(report["sections"]) == 7
    assert report["sections"][1]["heading"] == "二、研究人群与主题分布"
    assert all("table" not in section for section in report["sections"])
    assert report["sections"][-1]["heading"] == "参考文献"
    assert [item["number"] for item in report["sections"][-1]["references"]] == [1, 2]
    assert all(item["text"] for item in report["sections"][-1]["references"])
    assert "#727" not in json.dumps(report, ensure_ascii=False)
    assert "〔REF_" not in json.dumps(report, ensure_ascii=False)
    assert "[1, 2]" in report["sections"][0]["content"]
    assert validate_final_report(report, 2) == []


def test_reference_numbers_follow_first_narrative_appearance():
    payload = _payload()
    payload["narrative"]["technical_scope"]["content"] = "第二篇先出现〔REF_002〕，随后引用第一篇〔REF_001〕。"
    payload["narrative"] = dict(reversed(list(payload["narrative"].items())))
    papers = [dict(item, reference_key=f"REF_{index:03d}") for index, item in enumerate(_papers(), start=1)]

    numbers = build_reference_number_map(payload, papers)
    report = build_deterministic_report(payload, papers)

    assert numbers == {"REF_002": 1, "REF_001": 2}
    assert "第二篇先出现[1]，随后引用第一篇[2]" in report["sections"][0]["content"]
    assert [item["paper_id"] for item in report["sections"][-1]["references"]] == [2, 1]


def test_word_report_renders_no_tables_apa_references_and_centered_page_number():
    report = build_deterministic_report(_payload(), _papers())

    document = Document(BytesIO(generate_word_report(report)))

    assert len(document.tables) == 0
    footer_paragraph = document.sections[0].footer.paragraphs[0]
    assert footer_paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert "PAGE" in footer_paragraph._p.xml
    reference_heading_index = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text == "参考文献"
    )
    reference_paragraphs = [
        paragraph for paragraph in document.paragraphs[reference_heading_index + 1:]
        if paragraph.text.startswith("[")
    ]
    assert len(reference_paragraphs) == 2
    assert reference_paragraphs[0].text.startswith("[1] ")
    assert reference_paragraphs[0].paragraph_format.first_line_indent is not None
    assert any(run.italic for run in reference_paragraphs[0].runs)


def test_validation_names_missing_content_and_rejects_meta_language():
    payload = _payload()
    payload["narrative"]["conclusion"]["content"] = "现有论文未提供足够证据"

    errors = validate_narrative_payload(payload)

    assert "SECTION_CONTENT_MISSING:conclusion" in errors
    assert "META_PHRASE:未提供足够证据" in errors

    payload["narrative"]["conclusion"]["content"] = "证据不足，因此无法判断。" * 4
    errors = validate_narrative_payload(payload)
    assert "META_PHRASE:证据不足" in errors
    assert "META_PHRASE:无法判断" in errors


class _FakeLlm:
    def __init__(self):
        broken = _payload()
        del broken["narrative"]["conclusion"]
        self.responses = [json.dumps(broken, ensure_ascii=False), json.dumps(_payload(), ensure_ascii=False)]
        self.models = []

    def chat(self, model, messages, **kwargs):
        self.models.append(model)
        return self.responses.pop(0)


def test_missing_section_is_repaired_with_same_model():
    llm = _FakeLlm()

    report, repair_count = generate_report_json(
        llm,
        "example-chat-model",
        {"max_output_tokens": 16000, "thinking": False},
        "系统提示",
        _papers(),
        max_output_tokens=16000,
    )

    assert repair_count == 1
    assert llm.models == ["example-chat-model", "example-chat-model"]
    assert len(report["sections"]) == 7


@pytest.fixture
def governance_db(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'app.db'}")
    db.reset_engines()
    GovernanceBase.metadata.create_all(db.get_engine())
    yield
    db.reset_engines()


def test_identical_active_job_is_reused(governance_db, monkeypatch):
    submitted = []
    monkeypatch.setattr(summary_jobs, "submit_summary_job", submitted.append)
    actor = {"actor_id": "local:limited", "username": "limited", "role": "limited"}

    first = summary_jobs.create_summary_job([727, 643], profile="report", model="example-chat-model", actor=actor)
    second = summary_jobs.create_summary_job([643, 727], profile="report", model="example-chat-model", actor=actor)

    assert first["job_id"] == second["job_id"]
    assert second["reused"] is True
    assert submitted == [first["job_id"], first["job_id"]]


def test_summary_api_uses_database_queue_id(governance_db, monkeypatch):
    monkeypatch.setenv("JOBS_INLINE", "true")
    monkeypatch.setattr(
        summary_jobs,
        "generate_summary_artifact",
        lambda *_args, **_kwargs: (
            {"success": True, "file_id": "test-file", "expires_at": 2_000_000_000},
            {"sections": []},
            0,
        ),
    )
    actor = {"actor_id": "local:admin", "username": "admin", "role": "admin"}

    created = summary_jobs.create_summary_job(
        [101], profile="report", model="example-chat-model", actor=actor
    )
    observed = summary_jobs.get_summary_job(created["job_id"], actor)

    assert created["job_id"]
    assert observed["job_id"] == created["job_id"]
    assert observed["status"] == "succeeded"
    assert observed["result"]["file_id"] == "test-file"


def test_summary_worker_startup_leaves_recovery_to_database_queue(governance_db, monkeypatch):
    submitted = []
    monkeypatch.setattr(summary_jobs, "submit_summary_job", submitted.append)
    with db.get_session() as session:
        session.add(SummaryJob(
            job_id="recover-me",
            fingerprint="fingerprint",
            actor_id="local:limited",
            actor_json=json.dumps({"actor_id": "local:limited", "role": "limited"}),
            paper_ids_json="[727,643]",
            profile="report",
            model="example-chat-model",
            status="running",
            progress=55,
            stage="正在生成报告内容",
            created_at_utc=utc_now(),
            updated_at_utc=utc_now(),
        ))

    summary_jobs.start_summary_job_worker()

    with db.get_session() as session:
        job = session.query(SummaryJob).filter_by(job_id="recover-me").one()
        assert job.status == "running"
        assert job.stage == "正在生成报告内容"
    assert submitted == []
