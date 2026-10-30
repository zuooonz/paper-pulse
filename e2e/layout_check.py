#!/usr/bin/env python3
"""检查主要页面在桌面与移动视口下的纵向布局。"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urljoin


VIEWPORTS = (
    (1440, 900),
    (1280, 720),
    (768, 1024),
    (390, 844),
    (360, 640),
)
PAGE_ROUTES = (
    ("papers", "/papers"),
    ("workspace", "/workspace"),
    ("authors", "/authors"),
    ("stats", "/stats"),
    ("settings", "/settings"),
    ("videos", "/videos"),
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", required=True, help="前端地址，例如 http://localhost:3000")
    parser.add_argument("--api-url", required=True, help="API 地址，例如 http://127.0.0.1:8000")
    parser.add_argument("--user", help="登录用户名；未提供时使用 admin")
    parser.add_argument("--password", help="登录密码；未提供时读取 E2E_USERS 的 admin 项")
    parser.add_argument("--out", required=True, help="截图和 JSON 报告输出目录")
    return parser.parse_args()


def credentials(args: argparse.Namespace) -> tuple[str, str]:
    if bool(args.user) != bool(args.password):
        raise ValueError("--user 和 --password 必须同时提供")
    if args.user and args.password:
        return args.user, args.password

    raw = os.environ.get("E2E_USERS", "")
    if not raw:
        raise ValueError("请提供 --user/--password，或设置含 admin 密码的 E2E_USERS")
    try:
        users = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError("E2E_USERS 不是有效的 JSON 对象") from exc
    if not isinstance(users, dict) or not isinstance(users.get("admin"), str) or not users["admin"]:
        raise ValueError("E2E_USERS 必须包含有效的 admin 密码")
    return "admin", users["admin"]


def element_name(element: dict[str, Any]) -> str:
    return str(element.get("name") or "<unnamed>")[:160]


def inspect_layout(page: Any) -> dict[str, Any]:
    return page.evaluate(
        """() => {
          const root = document.documentElement;
          const scroller = document.scrollingElement || root;
          const visible = (element) => {
            const style = getComputedStyle(element);
            const rect = element.getBoundingClientRect();
            return style.display !== 'none' && style.visibility !== 'hidden' && rect.width > 0 && rect.height > 0;
          };
          const name = (element) => {
            const tag = element.tagName.toLowerCase();
            const slot = element.getAttribute('data-slot');
            const id = element.id ? `#${element.id}` : '';
            const classes = [...element.classList].slice(0, 3).join('.');
            return `${tag}${id}${classes ? `.${classes}` : ''}${slot ? `[${slot}]` : ''}`;
          };
          const bars = [...document.querySelectorAll('[data-bottom-bar]')]
            .filter(visible)
            .map((element) => {
              const rect = element.getBoundingClientRect();
              return {
                name: name(element),
                x: rect.x, y: rect.y, width: rect.width, height: rect.height,
                right: rect.right, bottom: rect.bottom,
                withinViewport: rect.left >= -1 && rect.top >= -1 &&
                  rect.right <= window.innerWidth + 1 && rect.bottom <= window.innerHeight + 1,
              };
            });
          const controls = new Set(['INPUT', 'TEXTAREA', 'SELECT']);
          const scrollables = [...document.querySelectorAll('body *')].filter((element) => {
            if (!visible(element) || controls.has(element.tagName)) return false;
            const overflow = getComputedStyle(element).overflowY;
            return ['auto', 'scroll', 'overlay'].includes(overflow) && element.scrollHeight > element.clientHeight + 1;
          });
          const nested = [];
          for (const child of scrollables) {
            let parent = child.parentElement;
            while (parent && parent !== document.body && parent !== root) {
              if (scrollables.includes(parent)) {
                nested.push({
                  ancestor: name(parent), descendant: name(child),
                  ancestorExtra: parent.scrollHeight - parent.clientHeight,
                  descendantExtra: child.scrollHeight - child.clientHeight,
                });
              }
              parent = parent.parentElement;
            }
          }
          return {
            viewport: { width: window.innerWidth, height: window.innerHeight },
            document: {
              scrollWidth: root.scrollWidth,
              innerWidth: window.innerWidth,
              scrollHeight: scroller.scrollHeight,
              innerHeight: window.innerHeight,
            },
            bottomBars: bars,
            nestedScrollContainers: nested,
          };
        }"""
    )


def check_snapshot(page: Any, state: str, viewport: str, report: dict[str, Any]) -> None:
    try:
        layout = inspect_layout(page)
    except Exception as exc:
        report["checks"].append({
            "state": state,
            "viewport": viewport,
            "status": "failed",
            "error": f"无法读取页面布局：{exc}",
        })
        return

    failures: list[str] = []
    dimensions = layout["document"]
    if dimensions["scrollWidth"] > dimensions["innerWidth"]:
        failures.append(
            f"页面横向溢出：scrollWidth={dimensions['scrollWidth']}，innerWidth={dimensions['innerWidth']}"
        )
    if dimensions["scrollHeight"] > dimensions["innerHeight"] + 1:
        failures.append(
            f"整页纵向滚动：scrollHeight={dimensions['scrollHeight']}，innerHeight={dimensions['innerHeight']}"
        )
    outside = [element_name(bar) for bar in layout["bottomBars"] if not bar["withinViewport"]]
    if outside:
        failures.append(f"底栏超出视口：{', '.join(outside)}")
    for pair in layout["nestedScrollContainers"]:
        failures.append(
            "嵌套纵向滚动："
            f"{pair['ancestor']}（余量 {pair['ancestorExtra']}px）包含 "
            f"{pair['descendant']}（余量 {pair['descendantExtra']}px）"
        )

    report["checks"].append({
        "state": state,
        "viewport": viewport,
        "status": "failed" if failures else "passed",
        "layout": layout,
        "failures": failures,
    })


def record_failure(report: dict[str, Any], state: str, viewport: str, detail: str) -> None:
    report["checks"].append({
        "state": state,
        "viewport": viewport,
        "status": "failed",
        "failures": [detail],
    })


def screenshot(page: Any, out_dir: Path, viewport: str, state: str) -> str:
    target = out_dir / viewport / f"{state}.png"
    target.parent.mkdir(parents=True, exist_ok=True)
    page.screenshot(path=str(target), full_page=False, animations="disabled", timeout=15000)
    return str(target)


def wait_for_app(page: Any) -> None:
    try:
        page.wait_for_load_state("networkidle", timeout=2500)
    except Exception:
        pass
    page.wait_for_timeout(300)


def login(page: Any, user: str, password: str) -> bool:
    page.get_by_label("用户名").fill(user)
    page.get_by_label("密码").fill(password)
    page.get_by_role("button", name="登录").click()
    try:
        page.wait_for_url(re.compile(r"^(?!.*?/login(?:[?#]|$)).*"), timeout=15000)
    except Exception:
        return False
    wait_for_app(page)
    return "/login" not in page.url


def open_paper_detail(page: Any, width: int) -> bool:
    if width < 768:
        target = page.locator('[class*="md:hidden space-y-2"] button').first
    else:
        target = page.locator("table tbody tr").first
    if target.count() == 0:
        return False
    try:
        target.click(timeout=5000)
        page.locator('[data-slot="dialog-content"]').last.wait_for(state="visible", timeout=8000)
        return True
    except Exception:
        return False


def _visible_tab(page: Any, label: str) -> Any | None:
    tabs = page.get_by_role("tab", name=label, exact=True)
    for index in range(tabs.count()):
        tab = tabs.nth(index)
        if tab.is_visible():
            return tab
    return None


def _create_layout_test_draft(api_request: Any, api_url: str, token: str) -> int:
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_heading("布局走查测试草稿", level=1)
    document.add_paragraph("此草稿由布局检查脚本通过上传接口创建。")
    payload = BytesIO()
    document.save(payload)
    response = api_request.post(
        urljoin(api_url + "/", "api/drafts/upload-word"),
        headers={"Authorization": f"Bearer {token}"},
        multipart={
            "file": {
                "name": "layout-check.docx",
                "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "buffer": payload.getvalue(),
            },
        },
    )
    if response.status == 403:
        try:
            users = json.loads(os.environ.get("E2E_USERS", "{}"))
            admin_password = users.get("admin") if isinstance(users, dict) else None
        except json.JSONDecodeError:
            admin_password = None
        if isinstance(admin_password, str) and admin_password:
            login_response = api_request.post(
                urljoin(api_url + "/", "api/auth/login"),
                data={"username": "admin", "password": admin_password},
            )
            if login_response.status == 200:
                admin_token = login_response.json().get("access_token")
                if admin_token:
                    response = api_request.post(
                        urljoin(api_url + "/", "api/drafts/upload-word"),
                        headers={"Authorization": f"Bearer {admin_token}"},
                        multipart={
                            "file": {
                                "name": "layout-check.docx",
                                "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                "buffer": payload.getvalue(),
                            },
                        },
                    )
    if response.status != 200:
        raise RuntimeError(f"通过草稿上传 API 创建走查草稿失败：HTTP {response.status}")
    result = response.json()
    return int(result["draft_id"])


def select_workspace_draft(page: Any, api_request: Any, api_url: str, token: str) -> bool:
    draft_selector = page.locator("aside button").filter(has_text=re.compile(r"Cite ID"))

    def open_first_draft() -> bool:
        draft = draft_selector.first
        if draft.count() == 0:
            return False
        try:
            draft.click(timeout=5000)
            wait_for_app(page)
            dialog = page.locator('[data-slot="dialog-content"]')
            if dialog.count() and dialog.last.is_visible():
                dialog.last.wait_for(state="visible", timeout=3000)
            return True
        except Exception:
            return False

    for label in ("草稿", "已审核", "已发布"):
        tab = _visible_tab(page, label)
        if tab is None:
            continue
        try:
            tab.click(timeout=5000)
            wait_for_app(page)
        except Exception:
            continue
        if open_first_draft():
            return True

    try:
        _create_layout_test_draft(api_request, api_url, token)
        page.reload(wait_until="domcontentloaded", timeout=45000)
        wait_for_app(page)
        draft_tab = _visible_tab(page, "草稿")
        if draft_tab is not None:
            draft_tab.click(timeout=5000)
            wait_for_app(page)
        draft = page.locator("aside button").filter(has_text=re.compile(r"Cite ID")).filter(
            has_text=re.compile(r"布局走查测试草稿")
        ).first
        if draft.count() == 0:
            draft = page.locator("aside button").filter(has_text=re.compile(r"Cite ID")).first
        if draft.count() == 0:
            return False
        draft.click(timeout=5000)
        wait_for_app(page)
        dialog = page.locator('[data-slot="dialog-content"]')
        if dialog.count() and dialog.last.is_visible():
            dialog.last.wait_for(state="visible", timeout=3000)
        return True
    except Exception:
        return False


def main() -> int:
    args = parse_args()
    try:
        user, password = credentials(args)
    except ValueError as exc:
        print(str(exc), file=sys.stderr)
        return 2

    base_url = args.base_url.rstrip("/")
    api_url = args.api_url.rstrip("/")
    out_dir = Path(args.out).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    report: dict[str, Any] = {
        "started_at": datetime.now(timezone.utc).isoformat(),
        "base_url": base_url,
        "api_url": api_url,
        "viewports": [{"width": width, "height": height} for width, height in VIEWPORTS],
        "checks": [],
    }

    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("缺少 Playwright Python 包，请先安装 playwright 并运行 playwright install chromium。", file=sys.stderr)
        return 2

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        api_request = playwright.request.new_context(timeout=15000)
        try:
            api_token: str | None = None
            try:
                response = api_request.post(
                    urljoin(api_url + "/", "api/auth/login"),
                    data={"username": user, "password": password},
                )
                report["api_login_status"] = response.status
                api_login_ok = response.status == 200
                if api_login_ok:
                    api_token = response.json().get("access_token")
            except Exception as exc:
                report["api_login_status"] = None
                report["api_login_error"] = str(exc)
                api_login_ok = False

            for width, height in VIEWPORTS:
                viewport_name = f"{width}x{height}"
                context = browser.new_context(
                    viewport={"width": width, "height": height},
                    device_scale_factor=1,
                    is_mobile=width < 768,
                    has_touch=width < 768,
                )
                page = context.new_page()
                try:
                    try:
                        page.goto(urljoin(base_url + "/", "login"), wait_until="domcontentloaded", timeout=45000)
                        wait_for_app(page)
                    except Exception:
                        pass
                    try:
                        path = screenshot(page, out_dir, viewport_name, "login")
                        report.setdefault("screenshots", []).append(path)
                        check_snapshot(page, "login", viewport_name, report)
                    except Exception as exc:
                        record_failure(report, "login", viewport_name, f"截图失败：{exc}")
                    try:
                        logged_in = login(page, user, password)
                    except Exception:
                        logged_in = False
                    if not logged_in:
                        record_failure(report, "login", viewport_name, "前端登录未完成")
                    if not api_login_ok:
                        record_failure(report, "login", viewport_name, "API 登录预检失败")
                        continue

                    for state, route in PAGE_ROUTES:
                        page.goto(urljoin(base_url + "/", route.lstrip("/")), wait_until="domcontentloaded", timeout=45000)
                        wait_for_app(page)
                        if state == "papers":
                            try:
                                path = screenshot(page, out_dir, viewport_name, "papers")
                                report.setdefault("screenshots", []).append(path)
                                check_snapshot(page, "papers", viewport_name, report)
                            except Exception as exc:
                                record_failure(report, "papers", viewport_name, f"截图失败：{exc}")
                            opened = open_paper_detail(page, width)
                            if not opened:
                                record_failure(report, "paper-detail-dialog", viewport_name, "没有可打开的论文详情")
                            try:
                                path = screenshot(page, out_dir, viewport_name, "paper-detail-dialog")
                                report.setdefault("screenshots", []).append(path)
                                check_snapshot(page, "paper-detail-dialog", viewport_name, report)
                            except Exception as exc:
                                record_failure(report, "paper-detail-dialog", viewport_name, f"截图失败：{exc}")
                            continue
                        if state == "workspace":
                            selected = bool(api_token) and select_workspace_draft(page, api_request, api_url, api_token)
                            if not selected:
                                record_failure(report, "workspace-selected-draft", viewport_name, "没有可选择的草稿")
                            screenshot_name = "workspace-selected-draft"
                            state_name = screenshot_name
                        else:
                            screenshot_name = state
                            state_name = state
                        try:
                            path = screenshot(page, out_dir, viewport_name, screenshot_name)
                            report.setdefault("screenshots", []).append(path)
                            check_snapshot(page, state_name, viewport_name, report)
                        except Exception as exc:
                            record_failure(report, state_name, viewport_name, f"截图失败：{exc}")
                except Exception as exc:
                    record_failure(report, "viewport", viewport_name, f"页面检查异常：{exc}")
                finally:
                    context.close()
        finally:
            api_request.dispose()
            browser.close()

    report["finished_at"] = datetime.now(timezone.utc).isoformat()
    report["summary"] = {
        "passed": sum(check["status"] == "passed" for check in report["checks"]),
        "failed": sum(check["status"] == "failed" for check in report["checks"]),
    }
    failures = [
        {"state": check.get("state"), "viewport": check.get("viewport"), "failures": check.get("failures", [])}
        for check in report["checks"] if check["status"] == "failed"
    ]
    report["failed_list"] = failures
    report_path = out_dir / "layout-report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(report_path), "summary": report["summary"], "failed_list": failures}, ensure_ascii=False, indent=2))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
